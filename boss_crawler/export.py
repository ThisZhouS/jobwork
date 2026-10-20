"""导出：把清洗后的数据写成 Excel 工作簿（岗位明细 / 公司汇总 / 数据质量）。

openpyxl 不可用时只提示，不影响 CSV/JSONL 产物。
"""
from __future__ import annotations

from pathlib import Path

JOB_COLUMNS = [
    ("company_name", "公司名称"),
    ("legal_name", "工商全称(来自公司介绍)"),
    ("is_anonymous", "是否脱敏"),
    ("job_title", "岗位"),
    ("salary", "薪资原文"),
    ("salary_unit", "薪资单位"),
    ("experience", "经验"),
    ("degree", "学历"),
    ("city", "城市"),
    ("district", "区县"),
    ("business_district", "商圈"),
    ("address_detail", "详细地址"),
    ("address_full", "完整地址"),
    ("address_source", "地址来源"),
    ("industry", "行业"),
    ("skills_all", "技能(结构化)"),
    ("required_years", "年限要求"),
    ("degree_required", "学历要求"),
    ("work_schedule", "工时"),
    ("seniority", "资深度"),
    ("duty_count", "职责条数"),
    ("req_count", "要求条数"),
    ("benefits_structured", "福利(结构化)"),
    ("jd_length", "JD字数"),
    ("url", "链接"),
]

COMPANY_COLUMNS = [
    ("company_name", "公司名称"),
    ("legal_name", "工商全称(来自公司介绍)"),
    ("is_anonymous", "是否脱敏"),
    ("industry", "行业"),
    ("scale", "规模"),
    ("stage", "融资阶段"),
    ("job_count", "岗位数"),
    ("cities", "城市"),
    ("addresses", "地址"),
    ("address_source", "地址来源"),
    ("salary_range", "月薪区间"),
    ("sample_titles", "岗位示例"),
]


def _salary_range(company: dict) -> str:
    low, high = company["salary_min"], company["salary_max"]
    if low is None or high is None or low < 1000 or high < 1000:
        return ""
    return f"{low / 1000:.0f}K-{high / 1000:.0f}K"


def job_rows(jobs: list[dict]) -> list[dict]:
    rows = []
    for job in jobs:
        company, workplace = job["company"], job["workplace"]
        rows.append({
            "company_name": company["name"],
            "legal_name": company["legal_name"],
            "is_anonymous": "是" if company["is_anonymous"] else "",
            "job_title": job["title"],
            "salary": job["salary"],
            "salary_unit": job["salary_unit"],
            "experience": job["experience"],
            "degree": job["degree"],
            "city": workplace["city"],
            "district": workplace["district"],
            "business_district": workplace["business_district"],
            "address_detail": workplace["detail"],
            "address_full": workplace["full"],
            "address_source": workplace["source"],
            "industry": company["industry"],
            "skills_all": "、".join(job.get("skills_all") or []),
            "required_years": _years_text(job),
            "degree_required": job.get("degree_required") or "",
            "work_schedule": job.get("work_schedule") or "",
            "seniority": job.get("seniority") or "",
            "duty_count": len(job.get("responsibilities") or []),
            "req_count": len(job.get("requirements") or []),
            "benefits_structured": "、".join(job.get("jd_benefits") or []),
            "jd_length": job.get("jd_length") or 0,
            "url": job["url"],
        })
    return [{k: ("" if v is None else v) for k, v in row.items()} for row in rows]


def _years_text(job: dict) -> str:
    """把 min_years / max_years 拼成可读文本；无要求写「不限」，缺数据留空。"""
    low = job.get("min_years")
    high = job.get("max_years")
    if low is None:
        return ""
    if low == 0:
        return "不限"
    if high:
        return f"{low}-{high} 年"
    return f"{low} 年以上"


def company_rows(companies: list[dict]) -> list[dict]:
    return [
        {
            "company_name": c["company_name"],
            "legal_name": c["legal_name"],
            "is_anonymous": "是" if c["is_anonymous"] else "",
            "industry": c["industry"],
            "scale": c["scale"],
            "stage": c["stage"],
            "job_count": c["job_count"],
            "cities": " / ".join(c["cities"]),
            "addresses": " | ".join(c["addresses"]),
            "address_source": c.get("address_source", ""),
            "salary_range": _salary_range(c),
            "sample_titles": " / ".join(c["sample_titles"]),
        }
        for c in companies
    ]


def write_workbook(path: Path, jobs: list[dict], companies: list[dict], report: dict) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        print("未安装 openpyxl，跳过 xlsx 导出（pip install openpyxl）")
        return False

    wb = Workbook()
    wb.remove(wb.active)

    header_fill = PatternFill("solid", fgColor="1F4E79")
    header_font = Font(bold=True, color="FFFFFF")

    def add_sheet(title: str, columns: list[tuple[str, str]], rows: list[dict]) -> None:
        ws = wb.create_sheet(title)
        ws.append([label for _, label in columns])
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        for row in rows:
            ws.append(["" if (value := row.get(key)) is None else value for key, _ in columns])
        widths = [len(label) * 2 + 4 for _, label in columns]
        for index, row in enumerate(rows, start=2):
            for col, (key, _) in enumerate(columns):
                value = row.get(key, "")
                if isinstance(value, str) and len(value) > widths[col] - 4:
                    widths[col] = min(60, len(value) + 4)
                if key in {"addresses", "sample_titles", "url"}:
                    ws.cell(row=index, column=col + 1).alignment = Alignment(wrap_text=False)
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col)].width = min(60, max(10, width))
        ws.freeze_panes = "A2"

    add_sheet("岗位明细", JOB_COLUMNS, job_rows(jobs))
    add_sheet("公司汇总", COMPANY_COLUMNS, company_rows(companies))

    quality = [
        {"指标": "职位数", "数值": report["jobs"]},
        {"指标": "公司数", "数值": report["companies"]},
        {"指标": "地址来自详情接口", "数值": report["address_detail"]},
        {"指标": "地址由城市/商圈派生", "数值": report["address_derived"]},
        {"指标": "地址缺失", "数值": report["address_missing"]},
        {"指标": "脱敏公司（某大型…）", "数值": report["anonymous_companies"]},
        {"指标": "公司名被截断", "数值": report["truncated_companies"]},
        {"指标": "补全工商全称", "数值": report["legal_name_recovered"]},
        {"指标": "有 JD 正文", "数值": report.get("with_jd", 0)},
        {"指标": "抽取出岗位职责", "数值": report.get("with_duties", 0)},
        {"指标": "抽取出任职要求", "数值": report.get("with_requirements", 0)},
        {"指标": "技能实体数（去重）", "数值": report.get("skill_total", 0)},
        {"指标": "岗位-技能关联数", "数值": report.get("skill_links", 0)},
        {"指标": "明确年限要求", "数值": report.get("with_years", 0)},
    ]
    add_sheet("数据质量", [("指标", "指标"), ("数值", "数值")], quality)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return True
