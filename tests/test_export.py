import tempfile
import unittest
from pathlib import Path

from boss_crawler.structuring import build_companies, build_report, clean_record
from boss_crawler.store import expand_inputs, merge_record
from boss_crawler.export import company_rows, job_rows, write_workbook


def make_job(**overrides):
    record = {
        "job_id": "j1",
        "title": "Python开发",
        "company": "凌策科技",
        "company_intro": "凌策科技有限公司成立于2015年。",
        "address": "成都锦江区航天科技大厦32楼",
        "city": "成都",
        "salary": "10-15K",
        "salary_min": 10000,
        "salary_max": 15000,
        "industry": "计算机软件",
        "url": "https://www.zhipin.com/job_detail/j1.html",
    }
    record.update(overrides)
    return clean_record(record)


class ExpandInputsTests(unittest.TestCase):
    def test_glob_skips_clean_artifacts_but_explicit_path_kept(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "clean").mkdir()
            (root / "jobs_a.jsonl").write_text("{}\n", encoding="utf-8")
            (root / "jobs_clean.jsonl").write_text("{}\n", encoding="utf-8")
            (root / "clean" / "enriched.jsonl").write_text("{}\n", encoding="utf-8")

            globbed = [p.name for p in expand_inputs([str(root / "*.jsonl")])]
            self.assertIn("jobs_a.jsonl", globbed)
            self.assertNotIn("jobs_clean.jsonl", globbed)

            explicit = [p.name for p in expand_inputs([str(root / "clean" / "enriched.jsonl")])]
            self.assertEqual(explicit, ["enriched.jsonl"])

            # 目录展开同样跳过清洗目录，避免二次清洗
            self.assertEqual(expand_inputs([str(root / "clean")]), [])


class MergeRecordTests(unittest.TestCase):
    def test_longer_address_wins_and_empty_fields_filled(self):
        base = {"job_id": "1", "address": "成都", "description": ""}
        extra = {"job_id": "1", "address": "成都锦江区航天科技大厦32楼", "description": "JD 正文"}
        merged = merge_record(base, extra)
        self.assertEqual(merged["address"], "成都锦江区航天科技大厦32楼")
        self.assertEqual(merged["description"], "JD 正文")


class ExportTests(unittest.TestCase):
    def test_job_rows_mapping(self):
        rows = job_rows([make_job()])
        self.assertEqual(rows[0]["company_name"], "凌策科技")
        self.assertEqual(rows[0]["address_full"], "成都锦江区航天科技大厦32楼")
        self.assertEqual(rows[0]["address_source"], "detail")
        self.assertEqual(rows[0]["is_anonymous"], "")

    def test_company_rows_salary_range_and_anonymous_flag(self):
        jobs = [make_job(), make_job(job_id="j2", company="某大型互联网公司",
                                     company_intro="", address="西安雁塔区飞阳广场221",
                                     city="西安")]
        companies = build_companies(jobs)
        rows = {r["company_name"]: r for r in company_rows(companies)}
        self.assertEqual(rows["凌策科技"]["salary_range"], "10K-15K")
        self.assertEqual(rows["某大型互联网公司"]["is_anonymous"], "是")

    def test_company_address_source_marks_derived(self):
        derived = make_job(job_id="j3", address="", district="武侯区", business_district="中和")
        self.assertEqual(build_companies([derived])[0]["address_source"], "list")
        detailed = make_job(job_id="j4")
        self.assertEqual(build_companies([detailed])[0]["address_source"], "detail")

    def test_company_rows_skip_daily_salary(self):
        jobs = [make_job(salary="100-200元/天", salary_min=100, salary_max=200)]
        rows = company_rows(build_companies(jobs))
        self.assertEqual(rows[0]["salary_range"], "")

    def test_write_workbook(self):
        jobs = [make_job()]
        companies = build_companies(jobs)
        report = build_report(jobs, companies)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "jobs.xlsx"
            self.assertTrue(write_workbook(path, jobs, companies, report))
            from openpyxl import load_workbook

            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["岗位明细", "公司汇总", "数据质量"])
            self.assertEqual(wb["岗位明细"].max_row, 2)
            self.assertEqual(wb["数据质量"].cell(row=2, column=1).value, "职位数")


if __name__ == "__main__":
    unittest.main()
